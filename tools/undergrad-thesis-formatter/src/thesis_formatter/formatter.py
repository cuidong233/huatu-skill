from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from .analyzer import BODY_CANDIDATES, HEADING_CANDIDATES


LEVEL_1_PATTERNS = (
    re.compile(r"^第[一二三四五六七八九十\d]+[章节]\s*"),
    re.compile(r"^(摘要|abstract|参考文献|致谢|附录)\s*$", re.IGNORECASE),
)
LEVEL_2_PATTERNS = (
    re.compile(r"^[一二三四五六七八九十]+、"),
    re.compile(r"^\d+\.\d+\s+"),
)
LEVEL_3_PATTERNS = (
    re.compile(r"^（[一二三四五六七八九十]+）"),
    re.compile(r"^\d+\.\d+\.\d+\s+"),
)


def _find_style(document: Document, names: tuple[str, ...]):
    for name in names:
        try:
            return document.styles[name]
        except KeyError:
            continue
    return None


def _copy_font(source, target) -> None:
    target.name = source.name
    target.size = source.size
    target.bold = source.bold
    target.italic = source.italic
    target.underline = source.underline
    target.color.rgb = source.color.rgb


def _copy_paragraph_format(source, target) -> None:
    target.alignment = source.alignment
    target.first_line_indent = source.first_line_indent
    target.left_indent = source.left_indent
    target.right_indent = source.right_indent
    target.line_spacing = source.line_spacing
    target.line_spacing_rule = source.line_spacing_rule
    target.space_before = source.space_before
    target.space_after = source.space_after
    target.keep_together = source.keep_together
    target.keep_with_next = source.keep_with_next
    target.page_break_before = source.page_break_before
    target.widow_control = source.widow_control


def _copy_style(source_style, target_style) -> None:
    _copy_font(source_style.font, target_style.font)
    _copy_paragraph_format(source_style.paragraph_format, target_style.paragraph_format)


def _copy_page_setup(template_doc: Document, target_doc: Document) -> None:
    source = template_doc.sections[0]
    for section in target_doc.sections:
        section.page_width = source.page_width
        section.page_height = source.page_height
        section.orientation = source.orientation
        section.top_margin = source.top_margin
        section.bottom_margin = source.bottom_margin
        section.left_margin = source.left_margin
        section.right_margin = source.right_margin
        section.header_distance = source.header_distance
        section.footer_distance = source.footer_distance


def _sync_known_styles(template_doc: Document, target_doc: Document) -> dict[str, str]:
    style_map = {}

    template_body = _find_style(template_doc, BODY_CANDIDATES)
    target_body = _find_style(target_doc, BODY_CANDIDATES)
    if template_body is not None and target_body is not None:
        _copy_style(template_body, target_body)
        style_map["body"] = target_body.name

    for level, names in HEADING_CANDIDATES.items():
        source = _find_style(template_doc, names)
        target = _find_style(target_doc, names)
        if source is not None and target is not None:
            _copy_style(source, target)
            style_map[f"heading_{level}"] = target.name

    return style_map


def _detect_heading_level(text: str) -> int | None:
    cleaned = text.strip()
    if not cleaned:
        return None
    if any(pattern.match(cleaned) for pattern in LEVEL_1_PATTERNS):
        return 1
    if any(pattern.match(cleaned) for pattern in LEVEL_2_PATTERNS):
        return 2
    if any(pattern.match(cleaned) for pattern in LEVEL_3_PATTERNS):
        return 3
    return None


def _style_for_level(target_doc: Document, level: int):
    return _find_style(target_doc, HEADING_CANDIDATES[level])


def apply_template_format(
    template_path: str | Path,
    thesis_path: str | Path,
    output_path: str | Path,
) -> dict[str, int | str | dict[str, str]]:
    template_doc = Document(template_path)
    target_doc = Document(thesis_path)

    _copy_page_setup(template_doc, target_doc)
    style_map = _sync_known_styles(template_doc, target_doc)

    formatted_headings = 0
    formatted_body = 0
    body_style = _find_style(target_doc, BODY_CANDIDATES)

    for paragraph in target_doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue

        level = _detect_heading_level(text)
        if level is not None:
            style = _style_for_level(target_doc, level)
            if style is not None:
                paragraph.style = style
                formatted_headings += 1
            continue

        if body_style is not None:
            paragraph.style = body_style
            formatted_body += 1

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    target_doc.save(output)

    return {
        "output": str(output),
        "formatted_headings": formatted_headings,
        "formatted_body_paragraphs": formatted_body,
        "style_map": style_map,
    }
